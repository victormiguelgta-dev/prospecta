"""
Exporta os leads para Excel (planilha formatada) e para JSON.
"""
import json
import re
from datetime import datetime

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import PASTA_PLANILHAS
from .db import NOMES_CLASSES, NOMES_ETAPAS
from .limpeza import NOMES_REDES, SITUACOES
from .oportunidades import NOMES_SERVICOS

CEP_REGEX = re.compile(r"\b\d{5}-?\d{3}\b")


def _cep(endereco):
    achado = CEP_REGEX.search(endereco or "")
    return achado.group().replace("-", "") if achado else "99999999"


def para_excel(leads: list, nome_base: str = "leads") -> str:
    linhas = []
    for l in leads:
        redes = l.get("redes") or {}
        linha = {
            "Nome": l["nome"],
            "Classe": NOMES_CLASSES.get(l.get("classe"), ""),
            "Nota de potencial": l.get("nota_potencial"),
            "O que vender": ", ".join(NOMES_SERVICOS.get(o["servico"], o["servico"]) for o in l.get("oportunidades") or []),
            "Serviço principal": NOMES_SERVICOS.get(l.get("servico_principal"), ""),
            "Etapa": NOMES_ETAPAS.get(l.get("etapa"), ""),
            "Estado": l.get("estado") or "",
            "Categoria": l.get("categoria") or "",
            "Telefone": l.get("telefone") or "",
            "WhatsApp": l.get("whatsapp") or "",
            "E-mail": l.get("email") or "",
            "Situação do site": SITUACOES.get(l.get("situacao_site"), ""),
            "Site": l.get("site") or "",
            "Saúde do Maps": l.get("saude_maps"),
            "Nota Google": l.get("nota") or "",
            "Avaliações": l.get("avaliacoes") or "",
        }
        for rede, nome in NOMES_REDES.items():
            linha[nome] = redes.get(rede, "")
        linha.update({
            "Posts no Instagram": l.get("instagram_posts") or "",
            "Seguidores no Instagram": l.get("instagram_seguidores") or "",
            "Endereço": l.get("endereco") or "",
            "Região": l.get("regiao") or "",
            "Próximo passo": l.get("proximo_passo") or "",
            "Data do próximo passo": l.get("data_proximo_passo") or "",
            "Maps": l.get("maps_url") or "",
        })
        linhas.append(linha)

    df = pd.DataFrame(linhas)
    if not df.empty:
        df["_cep"] = df["Endereço"].apply(_cep)
        df = df.sort_values(by=["Nota de potencial", "_cep"], ascending=[False, True]).drop(columns=["_cep"])

    arquivo = PASTA_PLANILHAS / f"{nome_base}_{datetime.now():%Y-%m-%d_%H%M%S}.xlsx"
    df.to_excel(arquivo, index=False)

    # Formatação visual (cabeçalho azul, linhas zebradas, colunas ajustadas)
    wb = openpyxl.load_workbook(arquivo)
    ws = wb.active
    ws.title = "Leads"
    cabecalho = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    zebra = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    borda = Border(*(Side(style="thin", color="D9D9D9"),) * 4)

    for cell in ws[1]:
        cell.fill = cabecalho
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, linha in enumerate(ws.iter_rows(min_row=2), start=2):
        for cell in linha:
            cell.font = Font(name="Calibri", size=10)
            cell.border = borda
            cell.alignment = Alignment(vertical="center")
            if i % 2 == 0:
                cell.fill = zebra
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for coluna in ws.columns:
        maior = max((len(str(c.value)) for c in coluna if c.value), default=0)
        ws.column_dimensions[get_column_letter(coluna[0].column)].width = min(max(maior + 3, 12), 55)
    wb.save(arquivo)
    return str(arquivo)


def para_json(leads: list, nome_base: str = "leads") -> str:
    arquivo = PASTA_PLANILHAS / f"{nome_base}_{datetime.now():%Y-%m-%d_%H%M%S}.json"
    arquivo.write_text(json.dumps(leads, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(arquivo)
